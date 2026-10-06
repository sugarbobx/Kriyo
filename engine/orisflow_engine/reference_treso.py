"""Lecture SEULE du dernier classeur de trésorerie existant, pour comparer un
comptage du jour à celui de la veille. N'écrit jamais dans ce classeur.

La recherche du fichier le plus récent reprend `trouver_tresorerie_recente` du
script B (date la plus récente lue dans le nom du fichier, sous-dossiers inclus,
fichiers `~$` ignorés).
"""

from __future__ import annotations

import glob
import os
import re
from datetime import date
from typing import Optional

from openpyxl import load_workbook

from .regles_agences import AGENCE_COLONNE

NOM_FEUILLE_SYNTHESE = ("Synthèse", "Synthese")
LIGNE_TOTAL_COMPTES = 16


def extraire_date_nom(chemin: str) -> date:
    nom = os.path.basename(chemin)
    correspondance = re.search(r"(\d{2})[-\s_](\d{2})[-\s_](\d{4})", nom)
    if correspondance:
        try:
            return date(int(correspondance.group(3)), int(correspondance.group(2)), int(correspondance.group(1)))
        except ValueError:
            pass
    return date.min


def trouver_classeur_recent(dossier: str, avant: Optional[date] = None) -> Optional[str]:
    """`avant` (optionnel) : n'accepte que les classeurs dont la date lue dans le nom est
    strictement antérieure à `avant`. Nécessaire pour la génération (voir `generation.py`) :
    sans ce filtre, si le dossier de référence contient déjà un classeur daté du jour qu'on
    s'apprête à générer (ou d'un jour postérieur — rattrapage, saisie en avance...), ce
    classeur serait pris comme son propre modèle, et son propre J-1 comme référence de
    lui-même. Bug corrigé le 01/10/2026 (CLAUDE.md), repéré après l'ajout d'un classeur du
    30/09 dans le dossier de référence le jour même où Orisflow générait pour le 30/09."""
    if not dossier or not os.path.isdir(dossier):
        return None
    fichiers = glob.glob(os.path.join(dossier, "**", "*.xlsx"), recursive=True)
    fichiers += glob.glob(os.path.join(dossier, "**", "*.xlsm"), recursive=True)
    fichiers = [f for f in fichiers if not os.path.basename(f).startswith("~$")]
    if avant is not None:
        fichiers = [f for f in fichiers if extraire_date_nom(f) < avant]
    if not fichiers:
        return None
    return max(fichiers, key=extraire_date_nom)


def lire_totaux_comptes_precedents(dossier_reference: str) -> dict:
    """Retourne {agence_cle: total_ligne16} du dernier classeur trouvé, en LECTURE SEULE.

    Résultat : {"chemin": str|None, "date": str|None, "totaux": {agence_cle: int}}.
    Un total manquant ou illisible est simplement absent du dictionnaire `totaux`
    (aucune exception : l'absence de référence ne doit jamais bloquer l'import).
    """
    chemin = trouver_classeur_recent(dossier_reference)
    if chemin is None:
        return {"chemin": None, "date": None, "totaux": {}}

    totaux: dict[str, int] = {}
    try:
        classeur = load_workbook(chemin, data_only=True, read_only=True)
        feuille = next((classeur[n] for n in NOM_FEUILLE_SYNTHESE if n in classeur.sheetnames), None)
        if feuille is not None:
            for agence_cle, colonne in AGENCE_COLONNE.items():
                valeur = feuille[f"{colonne}{LIGNE_TOTAL_COMPTES}"].value
                if isinstance(valeur, (int, float)):
                    totaux[agence_cle] = int(valeur)
        classeur.close()
    except Exception:
        # Un classeur illisible ne doit jamais empêcher l'import : on renvoie ce qui a pu être lu.
        pass

    jour = extraire_date_nom(chemin)
    return {
        "chemin": chemin,
        "date": jour.isoformat() if jour != date.min else None,
        "totaux": totaux,
    }

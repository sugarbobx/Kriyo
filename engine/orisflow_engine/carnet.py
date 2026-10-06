"""Carnet interne des soldes bancaires (décision du 03/10/2026, option A).

Orisflow consigne chaque jour le solde de chaque compte bancaire suivi, dans un petit fichier
JSON interne (dossier de travail, jamais partagé). Il sert le lendemain quand un relevé
manque : l'utilisateur peut alors choisir « utiliser la valeur de la veille ».

Structure : { "AAAA-MM-JJ": { "cca:12": 186412276, "afriland:65": 134381673, ... } }
Écriture atomique (fichier temporaire puis renommage), comme le bordereau.
"""

from __future__ import annotations

import json
import os
import re
from datetime import date
from typing import Optional

from openpyxl import load_workbook

from .reference_treso import NOM_FEUILLE_SYNTHESE, extraire_date_nom
from .regles_agences import AGENCE_COLONNE
from .regles_banques import BONS_DE_CAISSE, RELEVES_ATTENDUS

NOM_FICHIER = "soldes_bancaires.json"

# Lignes du classeur « Synthèse » pour les banques dont le carnet garde un historique par
# compte (celles listées dans RELEVES_ATTENDUS) ; préfixe de carnet correspondant.
_LIGNE_ET_PREFIXE_BANQUE = {"cca_bank": (28, "cca"), "afriland": (29, "afriland"), "bgfi": (30, "bgfi")}


def lire(chemin: str) -> dict[str, dict[str, int]]:
    """Retourne le carnet complet, ou un dictionnaire vide si absent ou illisible."""
    if not chemin or not os.path.isfile(chemin):
        return {}
    try:
        with open(chemin, "r", encoding="utf-8") as fichier:
            donnees = json.load(fichier)
    except (OSError, json.JSONDecodeError):
        return {}
    return donnees if isinstance(donnees, dict) else {}


def valeur_de_la_veille(carnet: dict[str, dict[str, int]], cle: str, jour: date) -> Optional[int]:
    """Dernier solde connu de `cle` strictement antérieur à `jour`."""
    anterieurs = sorted(d for d in carnet if d < jour.isoformat() and cle in carnet[d])
    return carnet[anterieurs[-1]][cle] if anterieurs else None


def enregistrer(chemin: str, jour: date, soldes: dict[str, int]) -> None:
    """Ajoute ou remplace les soldes du jour, sans toucher aux autres jours."""
    if not chemin or not soldes:
        return
    carnet = lire(chemin)
    carnet.setdefault(jour.isoformat(), {}).update(soldes)
    os.makedirs(os.path.dirname(chemin) or ".", exist_ok=True)
    temporaire = chemin + ".tmp"
    with open(temporaire, "w", encoding="utf-8") as fichier:
        json.dump(carnet, fichier, ensure_ascii=False, indent=2, sort_keys=True)
    os.replace(temporaire, chemin)


def _comptes_attendus(banque: str, agence: str) -> list[str]:
    """Identifiants de compte (clé RIB ou numéro BGFI) attendus pour (banque, agence),
    dans le même ordre trié que `generation._ecrire_banques` les additionne — c'est cet
    ordre qui permet de ré-associer, en relisant la formule, chaque terme à son compte."""
    return sorted(cle.split(":", 1)[1] for cle, (_, b, a) in RELEVES_ATTENDUS.items() if b == banque and a == agence)


def _termes_formule(valeur: object) -> Optional[list[int]]:
    """Décompose une cellule écrite par Orisflow (ex. `=510000000+151847746+199692276`,
    ou un entier simple) en la liste de ses termes. `None` si la cellule ne ressemble pas
    à une addition reconstituée (classeur d'une autre origine, ou vide)."""
    if isinstance(valeur, (int, float)):
        return [int(valeur)]
    if isinstance(valeur, str) and valeur.startswith("=") and re.fullmatch(r"[\d\s+]+", valeur[1:] or ""):
        return [int(terme) for terme in valeur[1:].split("+")]
    return None


def extraire_soldes_classeur(chemin_classeur: str) -> tuple[date, dict[str, int], list[str]]:
    """Reconstitue, depuis un classeur de trésorerie que l'utilisateur a validé comme
    correct, les soldes bancaires par compte qu'il décrit — pour alimenter rétroactivement
    le carnet (décision du 05/10/2026, après comparaison d'un classeur généré à un classeur
    de référence). N'importe une cellule que lorsqu'elle se décompose sans ambiguïté : les
    bons de caisse fixes connus en tête, puis exactement un terme par compte attendu, dans
    l'ordre trié par identifiant — sinon elle est ignorée et signalée (jamais deviné).

    Retourne (jour lu dans le nom du fichier, soldes {"cca:12": montant, ...}, avertissements).
    """
    jour = extraire_date_nom(chemin_classeur)
    classeur = load_workbook(chemin_classeur, data_only=False)
    nom_feuille = next((n for n in NOM_FEUILLE_SYNTHESE if n in classeur.sheetnames), classeur.sheetnames[0])
    feuille = classeur[nom_feuille]

    soldes: dict[str, int] = {}
    avertissements: list[str] = []
    for banque, (ligne, prefixe) in _LIGNE_ET_PREFIXE_BANQUE.items():
        for agence, colonne in AGENCE_COLONNE.items():
            comptes = _comptes_attendus(banque, agence)
            if not comptes:
                continue
            valeur = feuille[f"{colonne}{ligne}"].value
            if valeur in (None, 0):
                continue  # pas de relevé ce jour-là pour cette agence : rien à en tirer
            termes = _termes_formule(valeur)
            if termes is None:
                avertissements.append(f"{prefixe} — {agence} : cellule {colonne}{ligne} illisible ({valeur!r}), ignorée.")
                continue
            fixe = BONS_DE_CAISSE.get((banque, agence), [])
            if termes[: len(fixe)] != fixe:
                avertissements.append(
                    f"{prefixe} — {agence} : bons de caisse attendus {fixe} introuvables dans {termes}, ignorée."
                )
                continue
            variables = termes[len(fixe) :]
            if len(variables) != len(comptes):
                avertissements.append(
                    f"{prefixe} — {agence} : {len(variables)} montant(s) pour {len(comptes)} compte(s) connu(s), "
                    "impossible de les distinguer sans ambiguïté — ignorée."
                )
                continue
            for compte, montant in zip(comptes, variables):
                soldes[f"{prefixe}:{compte}"] = montant
    return jour, soldes, avertissements
